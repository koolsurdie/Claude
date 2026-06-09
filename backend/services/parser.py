import io
from pathlib import Path


def parse_file(filename: str, content: bytes) -> str:
    """Parse a PDF, DOCX, or PPTX file and return extracted text."""
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return _parse_pdf(content)
    elif ext in (".docx", ".doc"):
        return _parse_docx(content)
    elif ext in (".pptx", ".ppt"):
        return _parse_pptx(content)
    elif ext == ".txt":
        return content.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _parse_pdf(content: bytes) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=content, filetype="pdf")
    texts = []
    for page in doc:
        texts.append(page.get_text())
    doc.close()
    return "\n\n".join(texts)


def _parse_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _parse_pptx(content: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    slide_texts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slide_texts.append(f"--- Slide {i} ---\n" + "\n".join(texts))
    return "\n\n".join(slide_texts)
