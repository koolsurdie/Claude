import io
from typing import List, Any


def build_storyboard_docx(project_name: str, screens: List[Any]) -> bytes:
    """Build a Word document storyboard from screen objects."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title_para = doc.add_heading(f"Storyboard: {project_name}", level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    for screen in screens:
        # Screen header
        screen_num = getattr(screen, "screen_number", "?")
        screen_type = getattr(screen, "screen_type", "content")
        heading = doc.add_heading(
            f"Screen {screen_num} — {screen_type.replace('_', ' ').title()}", level=1
        )

        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"

        def add_row(label: str, value: str):
            row = table.add_row()
            label_cell = row.cells[0]
            value_cell = row.cells[1]
            label_cell.width = Inches(1.8)
            label_cell.paragraphs[0].add_run(label).bold = True
            value_cell.paragraphs[0].add_run(str(value) if value else "")

        add_row("Title (EN)", getattr(screen, "title_en", "") or "")
        add_row("Title (FR)", getattr(screen, "title_fr", "") or "")
        add_row("Narration (EN)", getattr(screen, "narration_en", "") or "")
        add_row("Narration (FR)", getattr(screen, "narration_fr", "") or "")
        add_row("Visual Direction", getattr(screen, "visual_direction", "") or "")
        add_row("Interaction Type", getattr(screen, "interaction_type", "") or "")
        add_row("Synthesia Note", getattr(screen, "synthesia_note", "") or "")
        add_row("Vyond Note", getattr(screen, "vyond_note", "") or "")
        add_row("Articulate Note", getattr(screen, "articulate_note", "") or "")

        quiz = getattr(screen, "quiz", None)
        if quiz and isinstance(quiz, dict):
            add_row("Quiz (EN)", quiz.get("question_en", ""))
            add_row("Quiz (FR)", quiz.get("question_fr", ""))
            options = quiz.get("options", [])
            answer_idx = quiz.get("answer", 0)
            for i, opt in enumerate(options):
                marker = " ✓" if i == answer_idx else ""
                add_row(f"  Option {i+1}", f"{opt}{marker}")

        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def build_quiz_xlsx(project_name: str, screens: List[Any]) -> bytes:
    """Build an Excel quiz bank from screens that have quiz data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Quiz Bank"

    orange = "FF7900"
    headers = [
        "Screen #", "Screen Type",
        "Question (EN)", "Question (FR)",
        "Option A", "Option B", "Option C", "Option D",
        "Correct Answer", "Notes"
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=orange)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for screen in screens:
        quiz = getattr(screen, "quiz", None)
        if not quiz or not isinstance(quiz, dict):
            continue

        options = quiz.get("options", [])
        answer_idx = quiz.get("answer", 0)
        answer_letter = ["A", "B", "C", "D"][answer_idx] if answer_idx < 4 else "A"

        row = [
            getattr(screen, "screen_number", ""),
            getattr(screen, "screen_type", ""),
            quiz.get("question_en", ""),
            quiz.get("question_fr", ""),
            options[0] if len(options) > 0 else "",
            options[1] if len(options) > 1 else "",
            options[2] if len(options) > 2 else "",
            options[3] if len(options) > 3 else "",
            answer_letter,
            getattr(screen, "articulate_note", "") or "",
        ]
        ws.append(row)

    # Column widths
    col_widths = [10, 14, 40, 40, 25, 25, 25, 25, 14, 30]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
